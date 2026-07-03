from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
from datetime import datetime, timedelta
import locale
import logging
import pandas as pd
import pytz
import re
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import NamedStyle
from openpyxl.utils import get_column_letter

CATEGORIAS = {
    "Uber": ["uber"],
    "Estacionamento": ["park", "estacionamento"],
    "Posto": ["posto", "shell", "ipiranga"],
    "Bagagem": ["gol"]
    # "Refeição": ["giraffas", "outback", "lanchonete", "restaurante"]
}

# Configuração de logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)

# Locale para datas e moeda
try:
    locale.setlocale(locale.LC_ALL, 'pt_BR.UTF-8')
except locale.Error:
    logging.warning("Locale pt_BR.UTF-8 não disponível neste sistema. Usando padrão.")


@dataclass
class Transacao:
    Data: datetime
    Lançamento: str
    Valor: float
    Categoria: str
    Observações: str | None = None


class FinanceiroProcessor:
    def __init__(self, base_dir: Path):
        self.base_dir = base_dir
        self.hoje = datetime.now(pytz.timezone("America/Sao_Paulo"))
        self.passado = self.hoje - timedelta(days=30)
        self.futuro = self.hoje + timedelta(days=30)
        self.pasta_mes = self.base_dir / self.hoje.strftime('%y-%m-%b').title()
        self._setup_dirs()

    def _setup_dirs(self):
        if not self.pasta_mes.exists():
            (self.pasta_mes / "NF").mkdir(parents=True)
            logging.info(f"Pasta {self.pasta_mes} criada!")
        else:
            logging.info(f"Pasta {self.pasta_mes} já existe.")

    @staticmethod
    def find_header_row(file_path: Path, expected_columns: list[str]) -> int | None:
        df = pd.read_excel(file_path, header=None)
        for i, row in df.iterrows():
            if all(col in row.values for col in expected_columns):
                return i
        return None

    def load_monthly_data(self, meses: list[str]) -> list[Transacao]:
        transacoes: list[Transacao] = []
        for mes in meses:
            file_path = (
                self.base_dir
                / self.passado.strftime('%y-%m-%b').title()
                / f"Fatura_{mes}.xlsx"
            )
            try:
                header_row = self.find_header_row(
                    file_path,
                    [
                        'Data da transação', 'Lançamentos', 'Valor', 'Moeda',
                        'Valor (US$)', 'Cotação/Dia (R$)', 'Valor em R$', 'Despesa no Exterior'
                    ]
                )
                if header_row is not None:
                    df = pd.read_excel(file_path, header=header_row)
                    df = self.clean_data(df)
                    for _, row in df.iterrows():
                        transacoes.append(
                            Transacao(
                                Data=row['Data da transação'],
                                Lançamento=row['Lançamentos'],
                                Valor=row['Valor'],
                                Categoria=row['Categoria'],
                                Observações=row.get('Observações', None)
                            )
                        )
                    logging.info(f"Fatura de {mes} carregada com sucesso.")
                else:
                    logging.warning(f"Cabeçalho não encontrado em {mes}.")
            except FileNotFoundError:
                logging.warning(f"Arquivo para {mes} não encontrado.")
                continue
        return transacoes

    @staticmethod
    def categorize(descricao: str) -> str:
        desc = descricao.lower()
        for categoria, padroes in CATEGORIAS.items():
            if any(re.search(p, desc) for p in padroes):
                return categoria
        return "Refeição"

    def clean_data(self, df: pd.DataFrame) -> pd.DataFrame:
        df = df.drop(columns=['Unnamed: 0'], errors='ignore')
        df = df.dropna(subset=['Moeda'])
        df['Data da transação'] = pd.to_datetime(df['Data da transação'], errors='coerce')
        df = df.astype({'Valor': 'float32'})
        df = df[df['Despesa no Exterior'] == "Não"]
        df = df[df['Data da transação'].dt.month == self.passado.month]
        df = df.drop(
            columns=['Valor (US$)', 'Cotação/Dia (R$)', 'Valor em R$', 'Despesa no Exterior', 'Moeda'],
            errors='ignore'
        )
        df.insert(2, "Categoria", df.apply(lambda row: self.categorize(row["Lançamentos"]), axis=1))
        df["Observações"] = None
        return df

    def process(self):
        meses = [self.hoje.strftime('%B').title(), self.futuro.strftime('%B').title()]
        transacoes = self.load_monthly_data(meses)

        # Converter lista de Transacao em DataFrame
        df_final = pd.DataFrame([t.__dict__ for t in transacoes])

        output_file = (
            self.base_dir
            / self.passado.strftime('%y-%m-%b').title()
            / f"df_concatenado_categorias_{self.passado.strftime('%y-%m-%b').title()}.xlsx"
        )
        df_final.to_excel(output_file, index=False, engine='openpyxl')

        # Ajustar formatação no Excel
        wb = load_workbook(output_file)
        ws = wb.active

        # Definir como tabela
        tab = Table(displayName="Transacoes", ref=ws.dimensions)
        style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True, showColumnStripes=False)
        tab.tableStyleInfo = style
        ws.add_table(tab)

        # Estilo de número contábil
        contab = NamedStyle(name="contab")
        contab.number_format = u'_-R$ * #,##0.00_-'
        
        # Formatar colunas (data abreviada e contábil)
        for col in ws.iter_cols(min_col=1, max_col=ws.max_column):
            if col[0].value == "Data":
                for cell in col[1:]:
                    cell.number_format = "DD/MM/YYYY"
            elif col[0].value == "Valor":
                for cell in col[1:]:
                    cell.style = contab
    

        # Ajustar largura automática das colunas
        for col in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = max_length + 2  # margem
            ws.column_dimensions[col_letter].width = adjusted_width

        wb.save(output_file)
        logging.info(f"Arquivo Excel gerado com sucesso em {output_file}.")


if __name__ == "__main__":
    base_dir = Path(r'C:/Users/vital/OneDrive/Documentos/Trabalho/Despesas')
    processor = FinanceiroProcessor(base_dir)
    processor.process()

